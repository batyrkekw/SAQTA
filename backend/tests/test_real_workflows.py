"""Production workflows against SQLite and, when configured, isolated PostgreSQL schemas."""

from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
import io
import json
import os
from uuid import uuid4
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select, text, func
from sqlalchemy.orm import sessionmaker
from app.core.config import get_settings, now_utc, today_local
from app.db.session import Base, build_engine, get_db
from app.main import app
from app.models import (
    AuthUser,
    ForecastObservation,
    InventorySnapshot,
    Product,
    SurplusListing,
)
from app.forecasting.selection import select_forecast

DATABASES = ["sqlite"] + (["postgresql"] if os.getenv("SAQTA_TEST_DATABASE_URL") else [])


@pytest.fixture(params=DATABASES)
def workspace(request, tmp_path, monkeypatch):
    monkeypatch.setattr(get_settings(), "workspace_mode", "real")
    monkeypatch.setattr(get_settings(), "forecast_requires_subscription", False)
    monkeypatch.setattr(get_settings(), "llm_base_url", "")
    admin = None
    if request.param == "postgresql":
        url = os.environ["SAQTA_TEST_DATABASE_URL"]
        schema = "test_" + uuid4().hex
        admin = create_engine(url)
        with admin.begin() as connection:
            connection.execute(text(f"CREATE SCHEMA {schema}"))
        engine = create_engine(url, connect_args={"options": f"-csearch_path={schema}"})
    else:
        engine = build_engine(f"sqlite:///{tmp_path / 'workflows.db'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(engine, expire_on_commit=False)

    def override():
        with factory() as db:
            try:
                yield db
            except Exception:
                db.rollback()
                raise

    app.dependency_overrides[get_db] = override
    with TestClient(app) as client:
        yield client, factory
    app.dependency_overrides.clear()
    engine.dispose()
    if admin:
        with admin.begin() as connection:
            connection.execute(text(f"DROP SCHEMA {schema} CASCADE"))
        admin.dispose()


def account(workspace, name="owner", role="consumer", approve=False):
    client, factory = workspace
    response = client.post(
        "/api/auth/register",
        json={
            "email": f"{name}@example.test",
            "password": "test-password-long",
            "bakery_name": name,
            "address": "Test pickup street",
            "requested_role": role,
        },
    )
    assert response.status_code == 201, response.text
    user = response.json()
    assert user["user"]["role"] == "consumer"
    if approve:
        with factory() as db:
            db.get(AuthUser, user["user"]["id"]).role = role
            db.commit()
    return user, {"Authorization": "Bearer " + user["token"]}


def listing(workspace, headers, **changes):
    client, _ = workspace
    location = client.get("/api/locations", headers=headers).json()[0]["id"]
    payload = {
        "name": "Fresh bread",
        "category": "Bread",
        "unit": "pcs",
        "quantity": 5,
        "original_price": 1000,
        "price": 500,
        "description": "Pickup today",
        "location_id": location,
        "pickup_start": (now_utc() - timedelta(minutes=1)).isoformat(),
        "pickup_deadline": (now_utc() + timedelta(hours=1)).isoformat(),
        "expires_at": (now_utc() + timedelta(hours=2)).isoformat(),
        "bulk_available": True,
        "stock_confirmed": True,
        **changes,
    }
    return client.post("/api/market/listings", headers=headers, json=payload)


def reservation(client, headers, offer, **changes):
    return client.post(
        "/api/market/orders",
        headers=headers,
        json={
            "listing_id": offer,
            "quantity": 2,
            "customer_name": "Test recipient",
            "idempotency_key": "request-" + uuid4().hex,
            **changes,
        },
    )


def profile():
    return {
        "business_type": "Bakery",
        "working_days": [0, 1, 2, 3, 4, 5, 6],
        "products": [
            {
                "name": "Bread",
                "category": "Bread",
                "unit": "pcs",
                "typical_sold": 12,
                "typical_produced": 15,
                "typical_remaining": 3,
            }
        ],
    }


def import_request(client, headers, content, options, extension="csv", commit=False):
    options = {
        "source_origin": "real",
        "mode": "observations",
        "create_products": True,
        "default_unit": "шт.",
        "column_mapping": {"date": "date", "product": "product", "sold": "sold"},
        **options,
    }
    response = client.post(
        "/api/imports/preview",
        headers=headers,
        files={"file": ("data." + extension, content)},
        data={"options": json.dumps(options)},
    )
    if commit and response.status_code == 200:
        return client.post(
            "/api/imports/commit",
            headers=headers,
            files={"file": ("data." + extension, content)},
            data={
                "options": json.dumps(options),
                "preview_token": response.json()["preview_token"],
                "confirm": "IMPORT SAQTA",
            },
        )
    return response


def test_public_empty_role_escalation_and_owner_rules(workspace):
    client, _ = workspace
    assert client.get("/api/market/listings").json() == []
    _, buyer = account(workspace, "buyer")
    assert listing(workspace, buyer).status_code == 403
    _, owner = account(workspace, "business", "business", True)
    response = listing(workspace, owner)
    assert response.status_code == 201, response.text
    offer = response.json()
    assert len(client.get("/api/market/listings").json()) == 1
    assert (
        client.patch(f"/api/market/listings/{offer['id']}", headers=buyer, json={"description": "stolen"}).status_code
        == 403
    )
    _, stranger = account(workspace, "stranger", "business", True)
    assert (
        client.patch(
            f"/api/market/listings/{offer['id']}", headers=stranger, json={"description": "stolen"}
        ).status_code
        == 404
    )
    assert (
        client.patch(f"/api/market/listings/{offer['id']}", headers=owner, json={"description": "Updated"}).status_code
        == 200
    )
    assert (
        client.post(
            "/api/auth/register",
            json={
                "email": "hacker@example.test",
                "password": "long-password",
                "bakery_name": "Hacker",
                "address": "abc",
                "role": "business",
            },
        ).status_code
        == 422
    )
    assert reservation(client, owner, offer["id"]).status_code == 409


def test_reserve_cancel_idempotency_and_history(workspace):
    client, factory = workspace
    _, owner = account(workspace, "seller", "business", True)
    _, buyer = account(workspace, "buyer")
    _, other = account(workspace, "other")
    offer = listing(workspace, owner).json()
    first = reservation(client, buyer, offer["id"], idempotency_key="stable-key-123")
    assert first.status_code == 201, first.text
    again = reservation(client, buyer, offer["id"], idempotency_key="stable-key-123")
    assert again.json()["id"] == first.json()["id"]
    assert client.get("/api/market/listings").json()[0]["available_quantity"] == 3
    assert reservation(client, buyer, offer["id"], quantity=4).status_code == 409
    oid = first.json()["id"]
    assert client.post(f"/api/market/orders/{oid}/collect", headers=buyer).status_code == 403
    assert client.post(f"/api/market/orders/{oid}/cancel", headers=other).status_code == 404
    assert client.post(f"/api/market/orders/{oid}/cancel", headers=buyer).status_code == 200
    assert client.post(f"/api/market/orders/{oid}/cancel", headers=buyer).status_code == 409
    assert client.get("/api/market/listings").json()[0]["available_quantity"] == 5
    orders = client.get("/api/market/orders", headers=buyer).json()
    assert orders[0]["events"][-1]["status"] == "cancelled"
    assert client.get("/api/market/orders", headers=other).json() == []
    with factory() as db:
        assert db.scalar(select(InventorySnapshot)).on_hand == 5


def test_bulk_approval_collection_and_expiry(workspace):
    client, factory = workspace
    _, seller = account(workspace, "seller", "business", True)
    _, charity = account(workspace, "charity", "charity", True)
    _, buyer = account(workspace, "buyer")
    offer = listing(workspace, seller, terms="free", price=0, original_price=None).json()
    assert offer["original_price_kzt"] is None
    assert reservation(client, buyer, offer["id"], kind="bulk").status_code == 403
    order = reservation(
        client, charity, offer["id"], kind="bulk", quantity=5, beneficiaries=20, logistics="Van required"
    )
    assert order.status_code == 201, order.text
    oid = order.json()["id"]
    assert order.json()["status"] == "pending"
    assert reservation(client, charity, offer["id"], kind="bulk").status_code == 409
    assert client.post(f"/api/market/orders/{oid}/collect", headers=seller).status_code == 409
    assert client.post(f"/api/market/orders/{oid}/approve", headers=seller).status_code == 200
    assert client.post(f"/api/market/orders/{oid}/collect", headers=seller).status_code == 200
    assert client.post(f"/api/market/orders/{oid}/collect", headers=seller).status_code == 409
    with factory() as db:
        assert db.scalar(select(InventorySnapshot)).on_hand == 0
    expiring = listing(workspace, seller).json()
    order = reservation(client, buyer, expiring["id"]).json()
    with factory() as db:
        db.get(SurplusListing, expiring["id"]).pickup_deadline = now_utc() - timedelta(minutes=1)
        db.commit()
    assert reservation(client, buyer, expiring["id"]).status_code == 409
    rows = client.get("/api/market/orders", headers=buyer).json()
    assert rows[0]["status"] == "cancelled"
    assert client.get("/api/market/listings").json() == []


def test_parallel_reservations_no_overselling(workspace):
    client, factory = workspace
    _, seller = account(workspace, "seller", "business", True)
    _, buyer = account(workspace, "buyer")
    offer = listing(workspace, seller, quantity=3).json()
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(lambda _: reservation(client, buyer, offer["id"], quantity=2), range(2)))
    assert sorted(r.status_code for r in responses) == [201, 409], [r.text for r in responses]
    with factory() as db:
        row = db.get(SurplusListing, offer["id"])
        assert row.available_quantity == 1 and row.reserved_quantity == 2


@pytest.mark.parametrize(
    "changes",
    [
        {"quantity": 0},
        {"price": 1001},
        {"pickup_deadline": (now_utc() - timedelta(minutes=1)).isoformat()},
        {"expires_at": (now_utc() - timedelta(minutes=1)).isoformat()},
        {"stock_confirmed": False},
    ],
)
def test_listing_validation(workspace, changes):
    client, _ = workspace
    _, seller = account(workspace, "seller", "business", True)
    assert listing(workspace, seller, **changes).status_code == 422
    assert client.get("/api/market/listings").json() == []


def test_photo_is_sanitized(workspace):
    client, _ = workspace
    _, seller = account(workspace, "seller", "business", True)
    offer = listing(workspace, seller).json()
    assert (
        client.post(
            f"/api/market/listings/{offer['id']}/photo", headers=seller, files={"file": ("evil.svg", b"<svg/>")}
        ).status_code
        == 422
    )
    from PIL import Image

    image = io.BytesIO()
    Image.new("RGB", (10, 10)).save(image, format="PNG")
    assert (
        client.post(
            f"/api/market/listings/{offer['id']}/photo", headers=seller, files={"file": ("photo.png", image.getvalue())}
        ).status_code
        == 200
    )
    photo = client.get(f"/api/market/listings/{offer['id']}/photo")
    assert photo.headers["content-type"] == "image/jpeg"


def test_questionnaire_export_reimport_no_fake_history(workspace):
    client, factory = workspace
    _, headers = account(workspace)
    _, foreign = account(workspace, "other")
    response = client.put("/api/analysis/profile", headers=headers, json=profile())
    assert response.status_code == 200, response.text
    estimate = client.post("/api/analysis/runs", headers=headers, json={}).json()
    assert estimate["items"][0]["status"] == "estimate" and estimate["items"][0]["demand"] == 84
    assert estimate["items"][0]["metrics"] is None
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(ForecastObservation)) == 0
    response = client.get("/api/analysis/profile.xlsx", headers=headers)
    assert response.status_code == 200
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Business", "Products", "HistoricalData", "Assumptions"]
    assert book["HistoricalData"].max_row == 1
    back = client.post(
        "/api/analysis/profile/inspect", headers=headers, files={"file": ("business.xlsx", response.content)}
    )
    assert back.status_code == 200, back.text
    assert back.json()["products"][0]["typical_sold"] == 12
    assert client.put("/api/analysis/profile", headers=headers, json=back.json()).status_code == 200
    assert client.get("/api/analysis/state", headers=foreign).json()["profile"] is None
    assert client.get("/api/analysis/profile.xlsx", headers=foreign).status_code == 404
    assert client.post(f"/api/analysis/runs/{estimate['run_id']}/explain", headers=foreign).status_code == 404
    assert client.get("/api/analysis/state", headers=headers).json()["latest"]["run_id"] == estimate["run_id"]
    again = client.post("/api/analysis/runs", headers=headers, json={}).json()
    assert again["run_id"] == estimate["run_id"] and again["cached"]


def test_unknown_profile_and_export_injection(workspace):
    client, _ = workspace
    _, headers = account(workspace)
    payload = profile()
    payload["products"][0]["name"] = '=WEBSERVICE("https://example.test")'
    payload["products"][0]["typical_sold"] = None
    assert client.put("/api/analysis/profile", headers=headers, json=payload).status_code == 200
    report = client.post("/api/analysis/runs", headers=headers, json={}).json()
    assert report["items"][0]["status"] == "insufficient"
    book = openpyxl.load_workbook(io.BytesIO(client.get("/api/analysis/profile.xlsx", headers=headers).content))
    assert book["Products"]["A2"].data_type == "s"


@pytest.mark.parametrize("encoding,sep", [("utf-8-sig", ","), ("cp1251", ";"), ("utf-16", "\t")])
def test_import_encodings_automatic_catalog_and_duplicates(workspace, encoding, sep):
    client, factory = workspace
    _, headers = account(workspace)
    rows = [sep.join(["date", "product", "sold"])] + [
        sep.join([str(today_local() - timedelta(days=i)), "Хлеб", str(10 + i)]) for i in range(3)
    ]
    content = "\n".join(rows).encode(encoding)
    preview = import_request(client, headers, content, {})
    assert preview.status_code == 200 and preview.json()["can_commit"], preview.text
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(Product)) == 0
    saved = import_request(client, headers, content, {}, commit=True)
    assert saved.status_code == 200, saved.text
    assert saved.json()["imported_count"] == 3
    duplicate = import_request(client, headers, content, {}).json()
    assert duplicate["duplicate_rows"] == 3 and not duplicate["can_commit"]
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(InventorySnapshot)) == 0
    report = client.post("/api/analysis/runs", headers=headers, json={}).json()
    assert report["items"][0]["status"] == "ready" and report["items"][0]["metrics"] is None


@pytest.mark.parametrize(
    "row",
    ["bad-date,Bread,2", "2026-10-20,Bread,no", "2026-10-20,Bread,-1", "2026-10-20,Bread,2.5", "2026-10-20,=1+1,2"],
)
def test_invalid_rows_are_never_partially_saved(workspace, row):
    client, factory = workspace
    _, headers = account(workspace)
    content = ("date,product,sold\n" + row).encode()
    preview = import_request(client, headers, content, {})
    assert preview.status_code == 200 and preview.json()["errors"], preview.text
    response = import_request(client, headers, content, {}, commit=True)
    assert response.status_code == 422
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(ForecastObservation)) == 0


def test_xlsx_kazakh_headers_multisheet_and_ambiguity(workspace):
    client, _ = workspace
    _, headers = account(workspace)
    book = openpyxl.Workbook()
    book.active.title = "Notes"
    book.active.append(["note"])
    book.active.append(["read me"])
    sheet = book.create_sheet("Sales")
    sheet.append(["күні", "тауар", "сатылды", "қалдық"])
    sheet.append([str(today_local()), "Нан", 5, 2])
    output = io.BytesIO()
    book.save(output)
    response = client.post(
        "/api/imports/inspect",
        headers=headers,
        files={"file": ("data.xlsx", output.getvalue())},
        data={"source_origin": "real"},
    ).json()
    sales = next(s for s in response["sheets"] if s["name"] == "Sales")
    assert sales["mapping"]["sold"] == "сатылды"
    assert (
        import_request(
            client, headers, output.getvalue(), {"sheet": "Sales", "column_mapping": sales["mapping"]}, "xlsx", True
        ).json()["imported_count"]
        == 1
    )
    from app.services.observation_import import observation_mapping

    assert "sold" not in observation_mapping(["sold", "продано"])


def test_daily_validation_tenant_isolation_and_real_metrics(workspace):
    client, factory = workspace
    _, headers = account(workspace)
    _, other = account(workspace, "other")
    client.put("/api/analysis/profile", headers=headers, json=profile())
    state = client.get("/api/analysis/state", headers=headers).json()
    pid = state["products"][0]["id"]
    lid = state["locations"][0]["id"]
    payload = {
        "product_id": pid,
        "location_id": lid,
        "date": str(today_local()),
        "sold": 10,
        "produced": 15,
        "remaining": 5,
        "opening": 0,
        "waste": 0,
    }
    assert client.post("/api/analysis/observations", headers=other, json=payload).status_code == 404
    assert (
        client.post("/api/analysis/observations", headers=headers, json={**payload, "remaining": 6}).status_code == 422
    )
    assert client.post("/api/analysis/observations", headers=headers, json=payload).status_code == 201
    assert client.post("/api/analysis/observations", headers=headers, json=payload).json()["duplicate"]
    assert (
        client.post(
            "/api/analysis/observations", headers=headers, json={**payload, "sold": 9, "remaining": 6}
        ).status_code
        == 409
    )
    rows = ["date,product,sold"] + [f"{today_local() - timedelta(days=i)},Bread,{10 + i % 7}" for i in range(1, 40)]
    assert import_request(client, headers, "\n".join(rows).encode(), {}, commit=True).status_code == 200
    report = client.post("/api/analysis/runs", headers=headers, json={}).json()
    item = report["items"][0]
    assert item["metrics"] and len(item["metrics"]["points"]) == 7
    assert item["validation_period"][1] < item["holdout_period"][0]
    assert client.get("/api/analysis/runs", headers=other).json() == []
    assert client.post("/api/analysis/runs", headers=other, json={"product_id": pid}).status_code == 404
    fallback = client.post(f"/api/analysis/runs/{report['run_id']}/explain", headers=headers).json()
    assert fallback["explanation_service"] == "not_configured"
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(InventorySnapshot)) == 0


def test_model_validation_cannot_see_holdout():
    history = [{"date": today_local() - timedelta(days=120 - i), "sales": 20 + (i % 7) * 3} for i in range(120)]
    first = select_forecast(history, today_local(), 7, None)
    changed = [{**r, "sales": r["sales"] if i < 113 else 10000} for i, r in enumerate(history)]
    second = select_forecast(changed, today_local(), 7, None)
    assert first["comparison"] == second["comparison"] and first["model"] == second["model"]
    assert second["metrics"]["mae"] > first["metrics"]["mae"]
    assert all(point["training_end"] < point["date"] for point in first["metrics"]["points"])
    zero = select_forecast([{**r, "sales": 0} for r in history], today_local(), 7, None)
    assert zero["metrics"]["wape"] is None and zero["metrics"]["mae"] == 0


def test_questionnaire_estimates_belong_to_one_location(workspace):
    client, _ = workspace
    _, headers = account(workspace)
    first = client.get("/api/locations", headers=headers).json()[0]["id"]
    second = client.post(
        "/api/locations", headers=headers, json={"name": "Second location", "address": "Second test address"}
    ).json()["id"]
    assert client.put("/api/analysis/profile", headers=headers, json=profile()).status_code == 422
    payload = {**profile(), "estimate_location_id": first}
    assert client.put("/api/analysis/profile", headers=headers, json=payload).status_code == 200
    report = client.post("/api/analysis/runs", headers=headers, json={}).json()
    assert {i["location_id"]: i["status"] for i in report["items"]} == {first: "estimate", second: "insufficient"}
    only_second = client.post("/api/analysis/runs", headers=headers, json={"location_id": second}).json()
    assert only_second["items"][0]["status"] == "insufficient"
    _, foreign = account(workspace, "foreign")
    foreign_id = client.get("/api/locations", headers=foreign).json()[0]["id"]
    assert (
        client.put(
            "/api/analysis/profile", headers=headers, json={**payload, "estimate_location_id": foreign_id}
        ).status_code
        == 422
    )


def test_new_catalog_units_are_explicit_and_import_quality_persists(workspace):
    from app.models import ImportBatch

    client, factory = workspace
    _, headers = account(workspace)
    content = f"date,product,sold\n{today_local()},Bread,4".encode()
    invalid = import_request(client, headers, content, {"default_unit": None}).json()
    assert not invalid["can_commit"] and "единицу" in invalid["errors"][0]["message"]
    saved = import_request(client, headers, content, {"default_unit": "pcs"}, commit=True)
    assert saved.status_code == 200
    with factory() as db:
        batch = db.get(ImportBatch, saved.json()["batch_id"])
        details = json.loads(batch.mapping_json)
        assert details["default_unit"] == "pcs" and details["quality"]["valid_rows"] == 1
        assert details["quality"]["warnings"]


def test_negotiated_terms_require_an_actual_agreed_price(workspace):
    client, _ = workspace
    _, seller = account(workspace, "seller", "business", True)
    _, charity = account(workspace, "charity", "charity", True)
    offer = listing(workspace, seller, terms="negotiated", price=0).json()
    assert client.get("/api/market/listings?min_discount=50").json() == []
    assert client.get("/api/market/listings?max_price=100").json() == []
    order = reservation(client, charity, offer["id"], kind="bulk", beneficiaries=10).json()
    assert not order["items"][0]["price_agreed"]
    url = f"/api/market/orders/{order['id']}/approve"
    assert client.post(url, headers=seller, json={}).status_code == 422
    assert client.post(url, headers=seller, json={"agreed_price_kzt": 1001}).status_code == 422
    confirmed = client.post(url, headers=seller, json={"agreed_price_kzt": 400})
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["total_kzt"] == 800 and confirmed.json()["items"][0]["price_agreed"]
    collected = client.post(f"/api/market/orders/{order['id']}/collect", headers=seller)
    assert collected.status_code == 200 and collected.json()["total_kzt"] == 800
