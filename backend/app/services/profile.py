"""Versioned estimates, spreadsheet round trip, and no fabricated history."""

import io
import json
from typing import Literal
import openpyxl
from fastapi import HTTPException
from pydantic import Field, model_validator, ValidationError
from sqlalchemy import select
from app.core.config import now_utc
from app.models import BusinessProfile, Product, ForecastObservation, BakeryLocation
from app.schemas import Input
from app.services.data_import import read_tables, key


class ProductEstimate(Input):
    name: str = Field(min_length=1, max_length=120)
    category: str = Field(default="Другое", min_length=1, max_length=60)
    unit: str = Field(default="шт.", min_length=1, max_length=20)
    price: int | None = Field(default=None, ge=0, le=10_000_000)
    typical_produced: int | None = Field(default=None, ge=0, le=1_000_000)
    typical_sold: int | None = Field(default=None, ge=0, le=1_000_000)
    typical_remaining: int | None = Field(default=None, ge=0, le=1_000_000)
    frequency: str = Field(default="", max_length=120)
    shelf_life_days: int | None = Field(default=None, ge=1, le=3650)


class ProfileInput(Input):
    schema_version: Literal[1] = 1
    business_type: str = Field(min_length=1, max_length=120)
    locations_count: int = Field(default=1, ge=1, le=1000)
    estimate_location_id: int | None = Field(default=None, gt=0)
    categories: str = Field(default="", max_length=500)
    working_days: list[int] = Field(default_factory=lambda: list(range(7)), min_length=1, max_length=7)
    hours: str = Field(default="", max_length=120)
    weekday_variation: str = Field(default="", max_length=500)
    products: list[ProductEstimate] = Field(min_length=1, max_length=100)
    demand: str = Field(default="", max_length=1000)
    weekend: str = Field(default="", max_length=1000)
    peak_hours: str = Field(default="", max_length=500)
    seasonality: str = Field(default="", max_length=1000)
    promotions: str = Field(default="", max_length=1000)
    leftover_causes: str = Field(default="", max_length=1000)
    holidays: str = Field(default="", max_length=1000)
    weather: str = Field(default="", max_length=1000)
    supply: str = Field(default="", max_length=1000)
    price_changes: str = Field(default="", max_length=1000)
    confidence: Literal["estimate", "unknown"] = "estimate"

    @model_validator(mode="after")
    def valid_lists(self):
        if len(set(self.working_days)) != len(self.working_days) or any(d < 0 or d > 6 for d in self.working_days):
            raise ValueError("Укажите уникальные дни недели от 0 до 6")
        if len({key(p.name) for p in self.products}) != len(self.products):
            raise ValueError("Названия товаров не должны повторяться")
        return self


def profile_json(db):
    row = db.scalar(select(BusinessProfile))
    return {**json.loads(row.answers_json), "updated_at": row.updated_at.isoformat()} if row else None


def save_profile(db, payload):
    owner = db.info["bakery_id"]
    from app.models import Bakery

    db.scalar(select(Bakery).where(Bakery.id == owner).with_for_update())
    row = db.scalar(select(BusinessProfile))
    if row is None:
        row = BusinessProfile(bakery_id=owner, schema_version=1, answers_json="{}")
        db.add(row)
    locations = list(db.scalars(select(BakeryLocation)))
    if payload.estimate_location_id is not None and payload.estimate_location_id not in {item.id for item in locations}:
        raise HTTPException(422, "Выберите свою точку для оценок анкеты")
    if payload.estimate_location_id is None:
        if len(locations) != 1:
            raise HTTPException(422, "Выберите свою точку для оценок анкеты")
        payload = payload.model_copy(update={"estimate_location_id": locations[0].id})
    row.answers_json = payload.model_dump_json()
    catalog = list(db.scalars(select(Product)))
    for item in payload.products:
        matches = [p for p in catalog if key(p.name) == key(item.name)]
        if len(matches) > 1:
            raise HTTPException(409, "В каталоге есть неоднозначные названия товаров")
        if not matches:
            db.add(
                Product(
                    bakery_id=owner,
                    name=item.name,
                    category=item.category,
                    unit=item.unit,
                    price_kzt=item.price or 0,
                    price_known=item.price is not None,
                )
            )
        elif key(matches[0].unit).rstrip(".") != key(item.unit).rstrip("."):
            raise HTTPException(
                422, "Единица не совпадает с каталогом; автоматическое преобразование единиц не выполняется"
            )
    db.commit()
    return profile_json(db)


def workbook_bytes(db):
    profile = profile_json(db)
    if not profile:
        raise HTTPException(404, "Сначала сохраните анкету")
    book = openpyxl.Workbook()
    business = book.active
    business.title = "Business"
    business.append(["field", "value"])
    business.append(["schema_version", 1])
    business.append(["created_at", now_utc().isoformat()])
    for field, value in profile.items():
        if field not in {"products", "updated_at", "schema_version"}:
            business.append([field, json.dumps(value, ensure_ascii=False)])
    products = book.create_sheet("Products")
    fields = list(ProductEstimate.model_fields)
    products.append(fields)
    for item in profile["products"]:
        products.append([item.get(field) for field in fields])
    history = book.create_sheet("HistoricalData")
    history.append(
        [
            "date",
            "product",
            "location",
            "unit",
            "sold",
            "remaining",
            "produced",
            "opening",
            "waste",
            "unit_price_kzt",
            "revenue",
        ]
    )
    for row in db.scalars(
        select(ForecastObservation)
        .where(ForecastObservation.source_origin == "real")
        .order_by(ForecastObservation.date)
    ):
        product, location = db.get(Product, row.product_id), db.get(BakeryLocation, row.location_id)
        history.append(
            [
                row.date.isoformat(),
                product.name,
                location.name,
                product.unit,
                row.sold,
                row.remaining,
                row.produced,
                row.opening,
                row.waste,
                row.unit_price_kzt,
                row.revenue,
            ]
        )
    assumptions = book.create_sheet("Assumptions")
    assumptions.append(["product", "field", "value", "source", "confidence"])
    for item in profile["products"]:
        for field in ("typical_produced", "typical_sold", "typical_remaining", "shelf_life_days"):
            assumptions.append(
                [
                    item["name"],
                    field,
                    item.get(field),
                    "questionnaire",
                    profile["confidence"] if item.get(field) is not None else "unknown",
                ]
            )
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
            cell.fill = openpyxl.styles.PatternFill("solid", fgColor="123C2F")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
            for cell in column:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Never emit user input as spreadsheet formulas.
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def read_profile_workbook(content, filename):
    tables = {table["name"]: table for table in read_tables(content, filename)}
    if not {"Business", "Products", "HistoricalData", "Assumptions"} <= tables.keys():
        raise HTTPException(422, "Нужна книга SAQTA с листами Business, Products, HistoricalData и Assumptions")
    try:
        data = {}
        for _, row in tables["Business"]["rows"]:
            field, value = row["field"], row["value"]
            if field == "created_at":
                continue
            data[field] = json.loads(value) if isinstance(value, str) else value
        data["products"] = [{k: v for k, v in row.items() if v is not None} for _, row in tables["Products"]["rows"]]
        # IDs are tenant-local; exported workbooks can be reviewed in a new account.
        data["estimate_location_id"] = None
        return ProfileInput.model_validate(data)
    except (ValueError, KeyError, TypeError, ValidationError):
        raise HTTPException(422, "Структура анкеты повреждена. Проверьте версию и значения книги")
